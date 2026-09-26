"""parsers/powerpoint.py -- .pptx"""

from __future__ import annotations

from pptx import Presentation

from . import imagenes as img
from .base import DocumentoIlegible, cerrar, documento_vacio

EXTENSIONES = {".pptx"}


def parse(ruta: str, nombre: str) -> dict:
    d = documento_vacio(nombre, ".pptx")
    d["formato"] = "PowerPoint"
    try:
        pres = Presentation(ruta)
    except Exception as exc:
        raise DocumentoIlegible(f"No se pudo abrir el .pptx: {exc}") from None

    d["n_paginas"] = len(pres.slides)
    candidatas = []

    for n, slide in enumerate(pres.slides, start=1):
        titulo_slide = ""
        try:
            if slide.shapes.title and slide.shapes.title.text.strip():
                titulo_slide = slide.shapes.title.text.strip()
                d["encabezados"].append(titulo_slide)
        except Exception:
            pass

        for forma in slide.shapes:
            # texto
            if forma.has_text_frame:
                for parrafo in forma.text_frame.paragraphs:
                    texto = "".join(r.text for r in parrafo.runs).strip()
                    if texto and texto != titulo_slide:
                        d["parrafos"].append(texto)
            # tablas
            if getattr(forma, "has_table", False):
                filas = [[c.text.strip() for c in f.cells] for f in forma.table.rows]
                filas = [f for f in filas if any(f)]
                if filas:
                    d["tablas"].append({"filas": filas, "n_filas": len(filas)})
            # imagenes: en un deck de arquitectura el diagrama ES el contenido
            if forma.shape_type is not None and getattr(forma, "image", None):
                try:
                    paquete = img.empaquetar(forma.image.blob, f"diapositiva {n}")
                    if paquete:
                        candidatas.append(paquete)
                except Exception:
                    pass

        # notas del presentador: ahi vive el porque de la lamina
        try:
            if slide.has_notes_slide:
                notas = slide.notes_slide.notes_text_frame.text.strip()
                if notas:
                    d["zonas"].append((f"notas diapositiva {n}", notas))
        except Exception:
            pass

    if d["encabezados"]:
        d["zonas"].append(("portada", "\n".join(d["encabezados"][:5])))

    d["imagenes"] = img.seleccionar(candidatas)

    props = pres.core_properties
    d["propiedades"] = {
        "titulo": props.title or "",
        "autor": props.author or "",
        "ultimo_modificado_por": props.last_modified_by or "",
        "creado": props.created.isoformat() if props.created else "",
        "modificado": props.modified.isoformat() if props.modified else "",
    }
    return cerrar(d)

"""parsers/excel.py -- .xlsx / .xlsm"""

from __future__ import annotations

import openpyxl

from .base import DocumentoIlegible, cerrar, documento_vacio

EXTENSIONES = {".xlsx", ".xlsm"}
MAX_FILAS = 200      # por hoja: suficiente para entender la estructura
MAX_COLUMNAS = 30
MAX_HOJAS = 15


def parse(ruta: str, nombre: str) -> dict:
    d = documento_vacio(nombre, ".xlsx")
    d["formato"] = "Excel"
    try:
        libro = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    except Exception as exc:
        raise DocumentoIlegible(f"No se pudo abrir el libro: {exc}") from None

    try:
        hojas = libro.sheetnames[:MAX_HOJAS]
        if len(libro.sheetnames) > MAX_HOJAS:
            d["avisos"].append(
                f"El libro tiene {len(libro.sheetnames)} hojas; se leyeron "
                f"las primeras {MAX_HOJAS}.")

        d["n_paginas"] = len(hojas)
        d["encabezados"] = list(hojas)

        for nombre_hoja in hojas:
            hoja = libro[nombre_hoja]
            filas = []
            for i, fila in enumerate(hoja.iter_rows(values_only=True)):
                if i >= MAX_FILAS:
                    d["avisos"].append(
                        f"La hoja '{nombre_hoja}' se trunco a {MAX_FILAS} filas.")
                    break
                celdas = ["" if v is None else str(v).strip()
                          for v in fila[:MAX_COLUMNAS]]
                if any(celdas):
                    filas.append(celdas)
            if filas:
                d["tablas"].append({
                    "filas": filas,
                    "n_filas": len(filas),
                    "hoja": nombre_hoja,
                })
                # la primera fila suele ser el encabezado: describe el dominio
                d["parrafos"].append(
                    f"Hoja '{nombre_hoja}': columnas " + ", ".join(
                        c for c in filas[0] if c))

        props = libro.properties
        d["propiedades"] = {
            "titulo": props.title or "",
            "autor": props.creator or "",
            "ultimo_modificado_por": props.lastModifiedBy or "",
            "creado": props.created.isoformat() if props.created else "",
            "modificado": props.modified.isoformat() if props.modified else "",
        }
    finally:
        libro.close()

    if d["encabezados"]:
        d["zonas"].append(("hojas", ", ".join(d["encabezados"])))

    return cerrar(d)
